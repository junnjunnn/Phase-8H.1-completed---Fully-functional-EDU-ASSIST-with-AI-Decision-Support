import hashlib
import json
import re
import zipfile
import zlib
from io import BytesIO
from xml.etree import ElementTree

from django.core import signing
from django.db import DatabaseError
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework import status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView


MAX_UPLOAD_SIZE = 2 * 1024 * 1024
MAX_WORKBOOK_SIZE = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_WORKSHEETS = 10
IMPORT_TOKEN_MAX_AGE = 30 * 60


class ImportFileError(Exception):
    pass


def _normalize_header(value):
    if not isinstance(value, str):
        return ''
    return re.sub(r'[^a-z0-9]', '', value.strip().lower())


def build_import_template(
    title,
    columns,
    instructions,
    dropdown_choices=None,
    reference_choices=None,
    column_guidance=None,
    example_rows=None,
    example_note='',
    date_fields=(),
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append([header for header, _, _ in columns])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{_excel_column(len(columns))}1'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    for index, (header, field_name, _) in enumerate(columns, start=1):
        sheet.column_dimensions[_excel_column(index)].width = max(16, min(len(header) + 4, 40))
        if field_name in {'order', 'capacity'}:
            for row_number in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row=row_number, column=index).number_format = '0'
        elif field_name in date_fields:
            for row_number in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row=row_number, column=index).number_format = 'yyyy-mm-dd'
        elif (
            field_name in {'username', 'email', 'employee_id', 'phone_number', 'code', 'grade_level_codes', 'strand_code', 'lrn', 'parent_contact'}
            or field_name.endswith(('_code', '_username'))
        ):
            for row_number in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row=row_number, column=index).number_format = '@'

    choices_by_field = dict(dropdown_choices or {})
    reference_guides = []
    headers_by_field = {
        field_name: header for header, field_name, _ in columns
    }
    for field_name, choices in (reference_choices or {}).items():
        values = []
        labels = []
        for choice in choices:
            if isinstance(choice, (tuple, list)) and len(choice) == 2:
                value, label = choice
                if value is None or not str(value).strip():
                    continue
                value = str(value).strip()
                values.append(value)
                if label is not None and str(label).strip():
                    labels.append(f'{value} = {str(label).strip()}')
            elif choice is not None and str(choice).strip():
                values.append(str(choice).strip())
        choices_by_field[field_name] = values
        if labels:
            reference_guides.append(
                f'{headers_by_field.get(field_name, field_name)} reference '
                f'(code = display name): {"; ".join(labels)}'
            )

    if choices_by_field:
        lookup_sheet = workbook.create_sheet('_Lookups')
        lookup_sheet.sheet_state = 'hidden'
        columns_by_field = {
            field_name: index
            for index, (_, field_name, _) in enumerate(columns, start=1)
        }
        lookup_column = 1
        for field_name, choices in choices_by_field.items():
            if field_name not in columns_by_field:
                continue
            values = list(dict.fromkeys(str(value).strip() for value in choices if value is not None and str(value).strip()))
            if not values:
                continue
            lookup_letter = _excel_column(lookup_column)
            for row_number, value in enumerate(values, start=1):
                cell = lookup_sheet.cell(row=row_number, column=lookup_column, value=value)
                cell.data_type = 's'
            list_name = f'ImportChoices{lookup_column}'
            workbook.defined_names.add(DefinedName(
                list_name,
                attr_text=f"'{lookup_sheet.title}'!${lookup_letter}$1:${lookup_letter}${len(values)}",
            ))
            validation = DataValidation(
                type='list',
                formula1=f'={list_name}',
                allow_blank=True,
                showErrorMessage=False,
            )
            sheet.add_data_validation(validation)
            target_column = _excel_column(columns_by_field[field_name])
            validation.add(f'{target_column}2:{target_column}{MAX_IMPORT_ROWS + 1}')
            lookup_column += 1

    guide = workbook.create_sheet('Instructions')
    guide.append([f'{title} bulk import instructions'])
    guide.append(['1. Add records only on the first worksheet, keeping its column names and order unchanged.'])
    guide.append(['2. Read the column guide below and use the current reference-code dropdowns where provided.'])
    guide.append(['3. Use the separate Example worksheet to understand a completed row; never import that worksheet.'])
    guide.append(['4. Save as .xlsx, upload it in EDU ASSIST, review every row, correct rejected rows, and preview again before importing.'])
    guide.append([f'Each upload supports up to {MAX_IMPORT_ROWS} data rows and a 2 MB file.'])
    guide.append([])
    guide.append(['Column', 'Required?', 'How to fill it'])
    guidance = column_guidance or {}
    for header, field_name, required in columns:
        guide.append([
            header,
            'Required' if required else 'Optional',
            guidance.get(field_name, 'Enter the value as text. Leave blank when optional.'),
        ])
    guide.append([])
    guide.append(['Important notes'])
    if any(field in {'grade_level_code', 'grade_level_codes'} for _, field, _ in columns):
        guide.append([
            'For a Grade Level Code column, enter the exact code shown in its current dropdown/reference, not the grade-level display name.'
        ])
    for instruction in instructions:
        guide.append([instruction])
    for reference_guide in reference_guides:
        guide.append([reference_guide])
    guide.column_dimensions['A'].width = 42
    guide.column_dimensions['B'].width = 16
    guide.column_dimensions['C'].width = 100
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    example = workbook.create_sheet('Example')
    example.append([f'EXAMPLE ONLY - {title.upper()} - DO NOT IMPORT THIS WORKSHEET'])
    example.append(['This worksheet is for guidance only. Copy only values you have verified and need onto the first worksheet.'])
    example.append([
        example_note
        or 'Sample values are illustrative; replace identifiers and optional details with the values required for your records.'
    ])
    example.append([header for header, _, _ in columns])
    for row in example_rows or ():
        example.append([row.get(field_name, '') for _, field_name, _ in columns])
    example.freeze_panes = 'A5'
    for index, (header, _, _) in enumerate(columns, start=1):
        example.column_dimensions[_excel_column(index)].width = max(16, min(len(header) + 4, 40))
    example.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(columns)))
    example['A1'].font = Font(bold=True, color='FFFFFF')
    example['A1'].fill = PatternFill(fill_type='solid', fgColor='C65911')
    example['A1'].alignment = Alignment(wrap_text=True)
    for cell in example[4]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    for row in example.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    return workbook


def _excel_column(index):
    name = ''
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def parse_import_workbook(
    upload,
    columns,
    required_fields,
    header_aliases,
    row_label,
    numeric_fields,
    allow_empty_rows=False,
    ignored_headers=(),
    import_sheet_name=None,
):
    if upload.size > MAX_UPLOAD_SIZE:
        raise ImportFileError('The selected file exceeds the 2 MB upload limit. Split the workbook into smaller files and try again.')
    if not (upload.name or '').lower().endswith('.xlsx'):
        raise ImportFileError('Select an Excel .xlsx file.')

    upload.seek(0)
    try:
        with zipfile.ZipFile(upload) as archive:
            entries = archive.infolist()
            if len(entries) > 1000:
                raise ImportFileError('The workbook contains too many internal files.')
            if sum(entry.file_size for entry in entries) > MAX_WORKBOOK_SIZE:
                raise ImportFileError('The uncompressed workbook is too large to process.')
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ImportFileError('Encrypted workbooks are not supported.')
            if any(
                entry.file_size > 10 * 1024 * 1024
                or (entry.compress_size and entry.file_size / entry.compress_size > 100)
                for entry in entries
            ):
                raise ImportFileError('The workbook contains an oversized or highly compressed file.')
            for entry in entries:
                if entry.filename.endswith(('.xml', '.rels')):
                    xml_content = archive.read(entry)
                    if b'<!doctype' in xml_content.lower() or b'<!entity' in xml_content.lower():
                        raise ImportFileError('The workbook contains unsupported XML declarations.')
            names = set(archive.namelist())
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                raise ImportFileError('The uploaded file is not a valid Excel workbook.')
    except ImportFileError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError, RuntimeError, EOFError, zlib.error) as exc:
        raise ImportFileError('The uploaded file is not a valid .xlsx workbook.') from exc
    finally:
        upload.seek(0)

    try:
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError, TypeError, IndexError, RuntimeError, EOFError, zlib.error, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ImportFileError('The uploaded file could not be read as an Excel workbook.') from exc

    try:
        if not workbook.worksheets or len(workbook.worksheets) > MAX_WORKSHEETS:
            raise ImportFileError('The workbook must contain between 1 and 10 worksheets.')
        sheet = next((
            item for item in workbook.worksheets
            if item.sheet_state == 'visible' and item.title == import_sheet_name
        ), None)
        if sheet is None:
            sheet = next((
                item for item in workbook.worksheets
                if item.sheet_state == 'visible' and item.title not in {'Instructions', 'Example'}
            ), None)
        if sheet is None:
            raise ImportFileError(f'The workbook must contain a visible {row_label.lower()} worksheet.')
        sheet.calculate_dimension(force=True)
        if sheet.max_row is None or sheet.max_row < 1:
            raise ImportFileError(f'The {row_label.lower()} worksheet is empty.')
        normalized_ignored_headers = {_normalize_header(header) for header in ignored_headers}
        if sheet.max_column is None or sheet.max_column > len(columns) + len(normalized_ignored_headers):
            raise ImportFileError(f'The {row_label.lower()} worksheet contains too many columns.')

        aliases = dict(header_aliases)
        for header, field_name, _ in columns:
            aliases.setdefault(_normalize_header(header), field_name)
            aliases.setdefault(_normalize_header(field_name), field_name)
        labels_by_field = {field_name: header for header, field_name, _ in columns}

        header_cells = next(sheet.iter_rows(min_row=1, max_row=1, max_col=sheet.max_column))
        headers = []
        for cell in header_cells:
            if cell.data_type == 'f':
                raise ImportFileError('Spreadsheet headers cannot contain formulas.')
            normalized_header = _normalize_header(cell.value)
            headers.append(
                '__ignored_legacy_header__'
                if normalized_header in normalized_ignored_headers
                else aliases.get(normalized_header, '')
            )
        if any(not header for header in headers):
            unsupported = [
                str(cell.value).strip() or f'column {_excel_column(index)}'
                for index, (cell, header) in enumerate(zip(header_cells, headers), start=1)
                if not header
            ]
            expected = ', '.join(header for header, _, _ in columns)
            raise ImportFileError(
                f'Unsupported or blank column name(s): {", ".join(unsupported)}. '
                f'Use the template headers: {expected}.'
            )
        if len(set(headers)) != len(headers):
            duplicate_labels = ', '.join(
                str(cell.value) for cell, field_name in zip(header_cells, headers)
                if headers.count(field_name) > 1
            )
            raise ImportFileError(
                f'Duplicate column(s): {duplicate_labels}. Keep each template column only once.'
            )
        missing_headers = required_fields - set(headers)
        if missing_headers:
            labels = ', '.join(
                header for header, field_name, _ in columns if field_name in missing_headers
            )
            raise ImportFileError(f'Missing required column(s): {labels}.')

        rows = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column),
            start=2,
        ):
            if all(cell.value is None for cell in cells):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ImportFileError(f'The workbook may contain at most {MAX_IMPORT_ROWS} {row_label.lower()} rows.')
            values = {}
            errors = []
            for field_name, cell in zip(headers, cells):
                if field_name == '__ignored_legacy_header__':
                    continue
                if cell.data_type == 'f':
                    errors.append(
                        f'Column "{labels_by_field[field_name]}" contains a formula; enter a value instead.'
                    )
                elif cell.value is None:
                    values[field_name] = ''
                elif field_name in numeric_fields and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    if float(cell.value).is_integer():
                        values[field_name] = str(int(cell.value))
                    else:
                        errors.append(
                            f'Column "{labels_by_field[field_name]}" has invalid value "{cell.value}"; '
                            'enter a whole number.'
                        )
                elif not isinstance(cell.value, str):
                    errors.append(
                        f'Column "{labels_by_field[field_name]}" has invalid value "{cell.value}"; '
                        'enter it as text.'
                    )
                else:
                    values[field_name] = (
                        cell.value
                        if field_name == 'official_subject_code'
                        else cell.value.strip()
                    )
            for field_name in required_fields:
                if field_name in values and not values[field_name]:
                    errors.append(f'Column "{labels_by_field[field_name]}" is required but is blank.')
            for _, field_name, _ in columns:
                values.setdefault(field_name, '')
            rows.append({'row': row_number, 'values': values, 'errors': errors})
        if not rows and not allow_empty_rows:
            raise ImportFileError(f'The worksheet does not contain any {row_label.lower()} rows.')
        return rows
    except ImportFileError:
        raise
    except (InvalidFileException, OSError, ValueError, KeyError, TypeError, IndexError, RuntimeError, EOFError, zlib.error, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ImportFileError(f'The {row_label.lower()} worksheet is malformed or could not be read.') from exc
    finally:
        workbook.close()


def preview_digest(parsed_rows):
    encoded = json.dumps(
        parsed_rows,
        sort_keys=True,
        separators=(',', ':'),
        ensure_ascii=True,
    ).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()


def add_row_error_context(rows, parsed_rows, columns):
    values_by_row = {row['row']: row['values'] for row in parsed_rows}
    labels = []
    for header, field_name, _ in columns:
        labels.append((header, field_name, re.sub(r'[^a-z0-9]', '', header.lower())))
        labels.append((header, field_name, re.sub(r'[^a-z0-9]', '', field_name.lower())))
    for row in rows:
        values = values_by_row.get(row['row'], {})
        contextual_errors = []
        for error in row.get('errors', []):
            normalized_error = re.sub(r'[^a-z0-9]', '', error.lower())
            match = next((
                (header, field_name)
                for header, field_name, normalized_header in sorted(
                    labels,
                    key=lambda item: len(item[2]),
                    reverse=True,
                )
                if (len(normalized_header) > 4 or normalized_header == 'lrn')
                and normalized_header in normalized_error
            ), None)
            if match:
                header, field_name = match
                value = values.get(field_name, '')
                if value and str(value).casefold() not in error.casefold():
                    error = f'{header} value "{value}": {error}'
                elif header.casefold() not in error.casefold():
                    error = f'{header}: {error}'
            elif not error.lower().startswith('row '):
                error = f'Row {row["row"]}: {error}'
            contextual_errors.append(error)
        row['errors'] = contextual_errors
    return rows


class BulkImportBase(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    columns = ()
    required_fields = frozenset()
    header_aliases = {}
    numeric_fields = frozenset()
    ignored_headers = ()
    import_title = 'Import'
    row_label = 'Record'
    token_salt = 'academic.bulk-import'
    filename = 'edu-assist-import-template.xlsx'
    allow_empty_rows = False

    def get(self, request):
        workbook = build_import_template(
            self.import_title,
            self.columns,
            self.template_instructions(),
            dropdown_choices=self.template_dropdown_choices(),
            reference_choices=self.template_reference_choices(),
            column_guidance=self.template_column_guidance(),
            example_rows=self.template_example_rows(),
            example_note=self.template_example_note(),
        )
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{self.filename}"'
        response['Cache-Control'] = 'no-store'
        return response

    def template_instructions(self):
        return []

    def template_dropdown_choices(self):
        return {}

    def template_reference_choices(self):
        return {}

    def template_column_guidance(self):
        return {}

    def template_example_rows(self):
        return []

    def template_example_note(self):
        return ''

    def parse(self, upload):
        return parse_import_workbook(
            upload,
            self.columns,
            self.required_fields,
            self.header_aliases,
            self.row_label,
            self.numeric_fields,
            allow_empty_rows=self.allow_empty_rows,
            ignored_headers=self.ignored_headers,
            import_sheet_name=self.import_title,
        )

    def review_rows(self, parsed_rows):
        raise NotImplementedError

    def persist_records(self, records, request):
        raise NotImplementedError

    def post(self, request):
        upload = request.FILES.get('file')
        if upload is None:
            return Response({'detail': 'Select an .xlsx file to preview.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            parsed_rows = self.parse(upload)
        except ImportFileError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        rows, valid_records = self.review_rows(parsed_rows)
        rows = add_row_error_context(rows, parsed_rows, self.columns)
        token = None
        if valid_records:
            token = signing.dumps(
                {'user_id': request.user.pk, 'digest': preview_digest(parsed_rows)},
                salt=self.token_salt,
                compress=True,
            )
        return Response({
            'rows': rows,
            'valid_count': len(valid_records),
            'rejected_count': len(rows) - len(valid_records),
            'import_token': token,
            'expires_in_seconds': IMPORT_TOKEN_MAX_AGE if token else 0,
        })


class BulkImportPreviewBase(BulkImportBase):
    http_method_names = ['post', 'options']


class BulkImportConfirmBase(BulkImportBase):
    http_method_names = ['post', 'options']
    importer_class = None

    def post(self, request):
        import_token = request.data.get('import_token')
        upload = request.FILES.get('file')
        if not isinstance(import_token, str) or upload is None:
            return Response({'detail': 'A valid preview and the same .xlsx workbook are required to confirm.'}, status=status.HTTP_400_BAD_REQUEST)
        importer = self.importer_class()
        try:
            parsed_rows = importer.parse(upload)
        except ImportFileError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        try:
            signed_preview = signing.loads(
                import_token,
                salt=importer.token_salt,
                max_age=IMPORT_TOKEN_MAX_AGE,
            )
        except signing.SignatureExpired:
            return Response({'detail': 'The preview has expired. Upload the workbook again.'}, status=status.HTTP_400_BAD_REQUEST)
        except signing.BadSignature:
            return Response({'detail': 'The preview could not be verified. Upload the workbook again.'}, status=status.HTTP_400_BAD_REQUEST)

        if (
            signed_preview.get('user_id') != request.user.pk
            or signed_preview.get('digest') != preview_digest(parsed_rows)
        ):
            return Response({'detail': 'The preview does not match this user or submitted data. Upload the workbook again.'}, status=status.HTTP_400_BAD_REQUEST)

        rows, valid_records = importer.review_rows(parsed_rows)
        rows = add_row_error_context(rows, parsed_rows, importer.columns)
        try:
            imported_rows = importer.persist_records(valid_records, request)
        except DatabaseError:
            return Response({
                'detail': 'The import was rolled back because a record changed or conflicted while saving. Preview the workbook again before retrying.',
            }, status=status.HTTP_409_CONFLICT)
        return Response({
            'imported_count': len(imported_rows),
            'imported_rows': imported_rows,
            'rejected_count': len(rows) - len(valid_records),
            'rejected_rows': [
                {'row': row['row'], 'errors': row['errors']}
                for row in rows if row['status'] != 'valid'
            ],
        })
