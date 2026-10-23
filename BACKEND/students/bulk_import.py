import datetime
import hashlib
import json
import re
import zipfile
import zlib
from collections import Counter
from io import BytesIO
from xml.etree import ElementTree

from django.core import signing
from django.db import DatabaseError, transaction
from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsRegistrarOrSchoolAdmin
from audit.models import AuditLog
from common.bulk_excel import add_row_error_context, build_import_template

from .models import Student
from .serializers import StudentSerializer


MAX_UPLOAD_SIZE = 2 * 1024 * 1024
MAX_WORKBOOK_SIZE = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_WORKSHEETS = 10
IMPORT_TOKEN_MAX_AGE = 30 * 60
IMPORT_TOKEN_SALT = 'students.bulk-registration'

TEMPLATE_COLUMNS = [
    ('LRN (Student ID)', 'lrn', True),
    ('First Name', 'first_name', True),
    ('Middle Name (Optional)', 'middle_name', False),
    ('Last Name', 'last_name', True),
    ('Suffix (Optional)', 'suffix', False),
    ('Gender (Male/Female)', 'gender', True),
    ('Birth Date (YYYY-MM-DD)', 'birth_date', True),
    ('Address', 'address', True),
    ('Guardian Name', 'guardian_name', True),
    ('Parent Contact', 'parent_contact', True),
    ('Email (Optional)', 'email', False),
]

HEADER_ALIASES = {
    'lrnstudentid': 'lrn',
    'lrn': 'lrn',
    'firstname': 'first_name',
    'middlenameoptional': 'middle_name',
    'middlename': 'middle_name',
    'lastname': 'last_name',
    'suffixoptional': 'suffix',
    'suffix': 'suffix',
    'gendermalefemale': 'gender',
    'gender': 'gender',
    'birthdateyyyymmdd': 'birth_date',
    'birthdate': 'birth_date',
    'address': 'address',
    'guardianname': 'guardian_name',
    'parentcontact': 'parent_contact',
    'emailoptional': 'email',
    'email': 'email',
}

REQUIRED_FIELDS = {
    'lrn',
    'first_name',
    'last_name',
    'gender',
    'birth_date',
    'address',
    'guardian_name',
    'parent_contact',
}
TEXT_FIELDS = {
    'lrn',
    'first_name',
    'middle_name',
    'last_name',
    'suffix',
    'gender',
    'address',
    'guardian_name',
    'parent_contact',
    'email',
}
FIELD_LABELS = {
    'lrn': 'LRN (Student ID)',
    'first_name': 'First name',
    'middle_name': 'Middle name',
    'last_name': 'Last name',
    'suffix': 'Suffix',
    'gender': 'Gender',
    'birth_date': 'Birth date',
    'address': 'Address',
    'guardian_name': 'Guardian name',
    'parent_contact': 'Parent contact',
    'email': 'Email',
}


class ImportFileError(Exception):
    pass


def build_student_template():
    candidate_lrns = [f'{suffix:012d}' for suffix in range(1, 100)]
    existing_lrns = set(
        Student.objects.filter(lrn__in=candidate_lrns).values_list('lrn', flat=True)
    )
    example_lrn = next(
        (lrn for lrn in candidate_lrns if lrn not in existing_lrns),
        '',
    )
    example_rows = [{
        'lrn': example_lrn,
        'first_name': 'Alex',
        'middle_name': 'Jordan',
        'last_name': 'Santos',
        'suffix': '',
        'gender': 'Male',
        'birth_date': '2012-04-15',
        'address': 'Example Street, Calinan',
        'guardian_name': 'Casey Santos',
        'parent_contact': '09123456789',
        'email': '',
    }] if example_lrn else []
    workbook = build_import_template(
        'Students',
        TEMPLATE_COLUMNS,
        [
            'Required: LRN, first name, last name, gender, birth date, address, guardian name, and parent contact. Middle name, suffix, and email are optional.',
            'New students are registered with Active status. This import registers student records only; enrollment into an academic year or section is separate.',
            'Enter LRN and parent contact as text to preserve leading zeroes. Numeric cells are rejected for those fields.',
            'Birth date accepts an Excel date cell, YYYY-MM-DD, or MM/DD/YYYY. Gender must be Male or Female.',
            'Student ID (LRN) must be unique. Do not change the header row.',
        ],
        dropdown_choices={'gender': ['Male', 'Female']},
        column_guidance={
            'lrn': 'Required unique student LRN/ID, entered as text to preserve leading zeroes; up to 12 characters. The importer checks that it is nonblank and unique, but does not validate a national ID checksum or numeric pattern.',
            'first_name': 'Required first name, up to 100 characters.',
            'middle_name': 'Optional middle name, up to 100 characters.',
            'last_name': 'Required last name, up to 100 characters.',
            'suffix': 'Optional name suffix, up to 20 characters.',
            'gender': 'Required. Choose Male or Female from the dropdown.',
            'birth_date': 'Required date. Use an Excel date cell, YYYY-MM-DD, or MM/DD/YYYY.',
            'address': 'Required address, up to 255 characters.',
            'guardian_name': 'Required guardian name, up to 150 characters.',
            'parent_contact': 'Required parent/guardian contact as text, up to 50 characters. Keep a leading zero by using Text cell format.',
            'email': 'Optional valid email address, up to 255 characters.',
        },
        example_rows=example_rows,
        example_note=(
            'The student ID and personal details below are fictitious format examples. Never register this sample row; replace it with authorized student data on the first worksheet.'
            if example_lrn else
            'No unused sample ID could be found in the current system data. Do not invent an LRN; use an authorized student LRN and preview the workbook.'
        ),
        date_fields={'birth_date'},
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _normalize_header(value):
    if not isinstance(value, str):
        return ''
    key = re.sub(r'[^a-z0-9]', '', value.strip().lower())
    return HEADER_ALIASES.get(key, '')


def _validate_archive(upload):
    upload.seek(0)
    try:
        with zipfile.ZipFile(upload) as archive:
            entries = archive.infolist()
            if len(entries) > 1000:
                raise ImportFileError('The workbook contains too many internal files.')
            total_size = sum(entry.file_size for entry in entries)
            if total_size > MAX_WORKBOOK_SIZE:
                raise ImportFileError('The uncompressed workbook is too large to process.')
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ImportFileError('Encrypted workbooks are not supported.')
            if any(
                entry.file_size > 10 * 1024 * 1024
                or (entry.compress_size and entry.file_size / entry.compress_size > 100)
                for entry in entries
            ):
                raise ImportFileError('The workbook contains an oversized or highly compressed file.')
            for entry in entries:
                if entry.filename.endswith(('.xml', '.rels')):
                    xml_content = archive.read(entry)
                    if b'<!doctype' in xml_content.lower() or b'<!entity' in xml_content.lower():
                        raise ImportFileError('The workbook contains unsupported XML declarations.')
            names = set(archive.namelist())
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                raise ImportFileError('The uploaded file is not a valid Excel workbook.')
    except (zipfile.BadZipFile, OSError, ValueError, RuntimeError, EOFError, zlib.error) as exc:
        raise ImportFileError('The uploaded file is not a valid .xlsx workbook.') from exc
    finally:
        upload.seek(0)


def _cell_text(value, field_name):
    if value is None:
        return ''
    if isinstance(value, bool):
        raise ValueError(f'{FIELD_LABELS[field_name]} must be text.')
    if isinstance(value, str):
        return value if field_name == 'parent_contact' else value.strip()
    if field_name in {'lrn', 'parent_contact'}:
        raise ValueError(f'{FIELD_LABELS[field_name]} value "{value}" must be entered as text. Set the Excel cell format to Text and re-enter it to preserve leading zeroes.')
    raise ValueError(f'{FIELD_LABELS[field_name]} value "{value}" must be entered as text.')


def _birth_date(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'Birth date value "{value}" is required and must be a valid date.')
    raw_value = value.strip()
    for date_format in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            return datetime.datetime.strptime(raw_value, date_format).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f'Birth date value "{raw_value}" is not accepted. Use an Excel date, YYYY-MM-DD, or MM/DD/YYYY.')


def _flatten_serializer_errors(errors):
    flattened = []
    for field_name, messages in errors.items():
        label = FIELD_LABELS.get(field_name, str(field_name).replace('_', ' ').capitalize())
        values = messages if isinstance(messages, (list, tuple)) else [messages]
        for message in values:
            flattened.append(f'{label}: {message}')
    return flattened


def _parse_workbook(upload):
    if upload.size > MAX_UPLOAD_SIZE:
        raise ImportFileError('The selected file exceeds the 2 MB upload limit. Split the workbook into smaller files and try again.')
    if not (upload.name or '').lower().endswith('.xlsx'):
        raise ImportFileError('Select an Excel .xlsx file.')

    _validate_archive(upload)
    try:
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError, TypeError, IndexError, RuntimeError, EOFError, zlib.error, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ImportFileError('The uploaded file could not be read as an Excel workbook.') from exc

    try:
        if not workbook.worksheets or len(workbook.worksheets) > MAX_WORKSHEETS:
            raise ImportFileError('The workbook must contain between 1 and 10 worksheets.')
        sheet = next((
            item for item in workbook.worksheets
            if item.sheet_state == 'visible' and item.title == 'Students'
        ), None)
        if sheet is None:
            sheet = next((
                item for item in workbook.worksheets
                if item.sheet_state == 'visible' and item.title not in {'Instructions', 'Example'}
            ), None)
        if sheet is None:
            raise ImportFileError('The workbook must contain a visible student worksheet.')
        sheet.calculate_dimension(force=True)
        if sheet.max_row is None or sheet.max_row < 1:
            raise ImportFileError('The student worksheet is empty.')
        if sheet.max_column is None or sheet.max_column > len(TEMPLATE_COLUMNS):
            raise ImportFileError('The student worksheet contains too many columns.')

        header_cells = next(sheet.iter_rows(min_row=1, max_row=1, max_col=sheet.max_column))
        headers = []
        for cell in header_cells:
            if cell.data_type == 'f':
                raise ImportFileError('Spreadsheet headers cannot contain formulas.')
            headers.append(_normalize_header(cell.value))
        if any(not header for header in headers):
            unsupported = [
                str(cell.value).strip() or f'column {index}'
                for index, (cell, header) in enumerate(zip(header_cells, headers), start=1)
                if not header
            ]
            expected = ', '.join(header for header, _, _ in TEMPLATE_COLUMNS)
            raise ImportFileError(
                f'Unsupported or blank column name(s): {", ".join(unsupported)}. '
                f'Use the template headers: {expected}.'
            )
        if len(set(headers)) != len(headers):
            duplicates = ', '.join(
                str(cell.value) for cell, header in zip(header_cells, headers)
                if headers.count(header) > 1
            )
            raise ImportFileError(
                f'Duplicate column(s): {duplicates}. Keep each template column only once.'
            )
        missing_headers = REQUIRED_FIELDS - set(headers)
        if missing_headers:
            labels = ', '.join(FIELD_LABELS[field] for field in sorted(missing_headers))
            raise ImportFileError(f'Missing required column(s): {labels}.')

        rows = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column),
            start=2,
        ):
            if all(cell.value is None for cell in cells):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ImportFileError(f'The workbook may contain at most {MAX_IMPORT_ROWS} student rows.')

            values = {}
            errors = []
            for header, cell in zip(headers, cells):
                if cell.data_type == 'f':
                    errors.append(f'Column "{FIELD_LABELS[header]}" contains a formula; enter a value instead.')
                    continue
                try:
                    if header == 'birth_date':
                        values[header] = _birth_date(cell.value) if cell.value is not None else ''
                    elif header in TEXT_FIELDS:
                        values[header] = _cell_text(cell.value, header)
                except ValueError as exc:
                    errors.append(str(exc))

            for field_name in REQUIRED_FIELDS:
                if field_name in values and not values[field_name]:
                    errors.append(f'{FIELD_LABELS[field_name]} is required.')
            if values.get('gender'):
                gender = values['gender'].casefold()
                if gender not in {'male', 'female'}:
                    errors.append(f'Gender value "{values["gender"]}" is not accepted. Enter Male or Female.')
                else:
                    values['gender'] = gender.title()

            for _, field_name, _ in TEMPLATE_COLUMNS:
                values.setdefault(field_name, '')

            rows.append({
                'row': row_number,
                'values': values,
                'errors': errors,
            })
        if not rows:
            raise ImportFileError('The worksheet does not contain any student rows.')
        return rows
    except ImportFileError:
        raise
    except (InvalidFileException, OSError, ValueError, KeyError, TypeError, IndexError, RuntimeError, EOFError, zlib.error, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ImportFileError('The student worksheet is malformed or could not be read.') from exc
    finally:
        workbook.close()


def _validated_payload(serializer):
    data = dict(serializer.validated_data)
    data.pop('current_enrollment', None)
    birth_date = data.get('birth_date')
    if birth_date:
        data['birth_date'] = birth_date.isoformat()
    return {
        field_name: data.get(field_name, '')
        for _, field_name, _ in TEMPLATE_COLUMNS
    }


def _review_rows(parsed_rows):
    lrn_counts = Counter(
        row['values'].get('lrn', '').strip()
        for row in parsed_rows
        if row['values'].get('lrn', '').strip()
    )
    existing_lrns = set(
        Student.objects.filter(lrn__in=lrn_counts).values_list('lrn', flat=True)
    )
    preview_rows = []
    valid_records = []

    for row in parsed_rows:
        errors = list(row['errors'])
        values = row['values']
        lrn = values.get('lrn', '').strip()
        duplicate_in_file = bool(lrn and lrn_counts[lrn] > 1)
        duplicate_in_database = bool(lrn and lrn in existing_lrns)
        if duplicate_in_file:
            errors.append('This LRN appears more than once in the uploaded spreadsheet.')
        if duplicate_in_database:
            errors.append('A student with this LRN is already registered.')

        serializer = None
        if not errors:
            serializer = StudentSerializer(data=values)
            if not serializer.is_valid():
                errors.extend(_flatten_serializer_errors(serializer.errors))

        is_valid = not errors
        preview_rows.append({
            'row': row['row'],
            'lrn': lrn,
            'first_name': values.get('first_name', ''),
            'last_name': values.get('last_name', ''),
            'status': (
                'duplicate_file' if duplicate_in_file
                else 'duplicate_existing' if duplicate_in_database
                else 'valid' if is_valid
                else 'invalid'
            ),
            'errors': errors,
        })
        if is_valid and serializer is not None:
            valid_records.append({
                'row': row['row'],
                'data': _validated_payload(serializer),
            })

    preview_rows = add_row_error_context(preview_rows, parsed_rows, TEMPLATE_COLUMNS)
    return preview_rows, valid_records


def _preview_digest(parsed_rows):
    preview = [
        {'row': row['row'], 'values': row['values'], 'errors': row['errors']}
        for row in parsed_rows
    ]
    encoded = json.dumps(preview, sort_keys=True, separators=(',', ':'), ensure_ascii=True).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()


class StudentImportTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsRegistrarOrSchoolAdmin]

    def get(self, request):
        response = HttpResponse(
            build_student_template().getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="edu-assist-student-import-template.xlsx"'
        response['Cache-Control'] = 'no-store'
        return response


class StudentImportPreviewView(APIView):
    permission_classes = [IsAuthenticated, IsRegistrarOrSchoolAdmin]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        upload = request.FILES.get('file')
        if upload is None:
            return Response({'detail': 'Select an .xlsx file to preview.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            parsed_rows = _parse_workbook(upload)
        except ImportFileError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        rows, valid_records = _review_rows(parsed_rows)
        import_token = None
        if valid_records:
            import_token = signing.dumps(
                {
                    'user_id': request.user.pk,
                    'digest': _preview_digest(parsed_rows),
                },
                salt=IMPORT_TOKEN_SALT,
                compress=True,
            )
        return Response({
            'rows': rows,
            'valid_count': len(valid_records),
            'rejected_count': len(rows) - len(valid_records),
            'import_token': import_token,
            'expires_in_seconds': IMPORT_TOKEN_MAX_AGE if import_token else 0,
        })


class StudentImportConfirmView(APIView):
    permission_classes = [IsAuthenticated, IsRegistrarOrSchoolAdmin]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        import_token = request.data.get('import_token')
        upload = request.FILES.get('file')
        if not isinstance(import_token, str) or upload is None:
            return Response({'detail': 'A valid preview and the same .xlsx workbook are required to confirm.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            parsed_rows = _parse_workbook(upload)
        except ImportFileError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        try:
            signed_preview = signing.loads(
                import_token,
                salt=IMPORT_TOKEN_SALT,
                max_age=IMPORT_TOKEN_MAX_AGE,
            )
        except signing.SignatureExpired:
            return Response({'detail': 'The preview has expired. Upload the workbook again.'}, status=status.HTTP_400_BAD_REQUEST)
        except signing.BadSignature:
            return Response({'detail': 'The preview could not be verified. Upload the workbook again.'}, status=status.HTTP_400_BAD_REQUEST)

        if (
            signed_preview.get('user_id') != request.user.pk
            or signed_preview.get('digest') != _preview_digest(parsed_rows)
        ):
            return Response({'detail': 'The preview does not match this user or submitted data. Upload the workbook again.'}, status=status.HTTP_400_BAD_REQUEST)

        rows, valid_records = _review_rows(parsed_rows)
        row_errors = {
            row['row']: row['errors']
            for row in rows
            if row['status'] != 'valid'
        }
        imported_rows = []
        if valid_records:
            try:
                with transaction.atomic():
                    for record in valid_records:
                        row_number = record['row']
                        serializer = StudentSerializer(data=record['data'])
                        if not serializer.is_valid():
                            row_errors[row_number] = _flatten_serializer_errors(serializer.errors)
                            continue
                        student = serializer.save()
                        AuditLog.objects.create(
                            user=request.user,
                            action='CREATE',
                            module='students',
                            object_type='Student',
                            object_id=str(student.pk),
                        )
                        imported_rows.append(row_number)
            except DatabaseError:
                imported_rows = []
                rollback_message = 'The import was rolled back because one or more records could not be saved. Preview the workbook again before retrying.'
                for record in valid_records:
                    row_number = record['row']
                    row_errors[row_number] = [rollback_message]

        rejected_rows = [
            {'row': row_number, 'errors': errors}
            for row_number, errors in sorted(row_errors.items())
        ]
        rejected_rows.sort(key=lambda row: row['row'])
        return Response({
            'imported_count': len(imported_rows),
            'imported_rows': imported_rows,
            'rejected_count': len(rejected_rows),
            'rejected_rows': rejected_rows,
        })
